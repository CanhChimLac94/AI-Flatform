"""
Build LLM-ready content parts from file attachments.

Returns OpenAI-compatible vision format:
  - Plain string when no images are present (all providers accept this).
  - List of content parts [{type: text/image_url}] when images are present.

_stream_anthropic converts image_url parts to Anthropic's native image format.
"""

import base64
import io
from pathlib import Path
from typing import Any

from app.schemas.chat import AttachmentRef

UPLOADS_DIR = Path(__file__).resolve().parent.parent.parent / "uploads"

IMAGE_CONTENT_TYPES = {"image/jpeg", "image/png", "image/gif", "image/webp"}


def _find_file(attachment_id: str) -> tuple[Path, bytes] | None:
    """Locate an uploaded file by its UUID and return (path, bytes)."""
    # Sanitise to prevent path traversal
    if not all(c.isalnum() or c == "-" for c in attachment_id):
        return None
    matches = list(UPLOADS_DIR.glob(f"{attachment_id}.*"))
    if not matches:
        return None
    path = matches[0]
    return path, path.read_bytes()


def _extract_pdf_text(data: bytes) -> str:
    try:
        import pypdf  # type: ignore
        reader = pypdf.PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(p for p in pages if p.strip())
    except Exception:
        return ""


def _extract_docx_text(data: bytes) -> str:
    try:
        import docx  # type: ignore
        doc = docx.Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception:
        return ""


def _extract_xlsx_text(data: bytes) -> str:
    try:
        import openpyxl  # type: ignore
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows: list[str] = []
        for sheet in wb.worksheets:
            rows.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                if any(c.strip() for c in cells):
                    rows.append("\t".join(cells))
        return "\n".join(rows)
    except Exception:
        return ""


def build_message_content(
    text: str,
    attachments: list[AttachmentRef],
) -> str | list[dict[str, Any]]:
    """
    Convert message text + attachments into LLM content.

    Returns:
      - str  — when no images are present (all providers accept plain string)
      - list — OpenAI vision content parts when images are present
    """
    if not attachments:
        return text

    doc_text_parts: list[str] = []
    image_parts: list[dict[str, Any]] = []

    for att in attachments:
        result = _find_file(att.id)
        if result is None:
            doc_text_parts.append(f"[Attachment not found: {att.name}]")
            continue
        path, data = result
        ct = att.content_type
        ext = path.suffix.lower()

        if ct in IMAGE_CONTENT_TYPES:
            b64 = base64.b64encode(data).decode()
            image_parts.append({
                "type": "image_url",
                "image_url": {"url": f"data:{ct};base64,{b64}"},
            })

        elif ct == "application/pdf" or ext == ".pdf":
            extracted = _extract_pdf_text(data)
            label = f"[File: {att.name}]"
            doc_text_parts.append(
                f"{label}\n{extracted}" if extracted
                else f"{label} — could not extract text from this PDF."
            )

        elif ct == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" or ext == ".docx":
            extracted = _extract_docx_text(data)
            label = f"[File: {att.name}]"
            doc_text_parts.append(
                f"{label}\n{extracted}" if extracted
                else f"{label} — could not extract text from this Word document."
            )

        elif ct == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" or ext == ".xlsx":
            extracted = _extract_xlsx_text(data)
            label = f"[File: {att.name}]"
            doc_text_parts.append(
                f"{label}\n{extracted}" if extracted
                else f"{label} — could not extract data from this spreadsheet."
            )

        else:
            doc_text_parts.append(f"[Unsupported file type: {att.name} ({ct})]")

    combined_text = text
    if doc_text_parts:
        combined_text = text + "\n\n" + "\n\n".join(doc_text_parts)

    if not image_parts:
        # Plain string — fully backwards-compatible with every provider
        return combined_text

    # Vision format — OpenAI-compatible list
    parts: list[dict[str, Any]] = [{"type": "text", "text": combined_text}]
    parts.extend(image_parts)
    return parts


def openai_to_anthropic_content(
    content: str | list[dict[str, Any]],
) -> str | list[dict[str, Any]]:
    """
    Convert OpenAI-compatible content to Anthropic's Messages API format.
    Plain strings are returned unchanged.
    image_url parts (data URI only) become Anthropic base64 image blocks.
    """
    if isinstance(content, str):
        return content

    anthropic_parts: list[dict[str, Any]] = []
    for part in content:
        if part["type"] == "text":
            anthropic_parts.append({"type": "text", "text": part["text"]})
        elif part["type"] == "image_url":
            url: str = part["image_url"]["url"]
            if url.startswith("data:"):
                # data:<media_type>;base64,<data>
                header, b64_data = url.split(",", 1)
                media_type = header.split(":")[1].split(";")[0]
                anthropic_parts.append({
                    "type": "image",
                    "source": {
                        "type": "base64",
                        "media_type": media_type,
                        "data": b64_data,
                    },
                })
    return anthropic_parts or content
